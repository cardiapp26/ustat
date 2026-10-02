"""The output document as one Word file.

The browser turns each output item (sanitised HTML, see
frontend/src/lib/outputBlocks.ts) into a flat list of blocks: headings,
paragraphs, tables as rows of cell text, and figures as base64 PNG/JPEG.
This module lays them out with python-docx. It does no statistics and
interprets no HTML: anything that is not one of the four block types is
rejected by the request model before it gets here.
"""
from __future__ import annotations

import base64
import binascii
import io
from typing import List, Literal, Optional

from pydantic import BaseModel, Field

# python-docx is imported inside the builder: routers/pub_export.py treats it
# as optional (HAS_DOCX), and a top-level import here would take that whole
# router down on an install without it.

MAX_ITEMS = 300
MAX_IMAGE_BYTES = 15 * 1024 * 1024
MAX_TABLE_CELLS = 20_000
_IMAGE_PREFIXES = ("data:image/png;base64,", "data:image/jpeg;base64,")
# Letter/A4 with default margins leaves ~6.3 in of text width.
_FIGURE_WIDTH_IN = 6.0


class OutputBlock(BaseModel):
    type: Literal["heading", "paragraph", "table", "image"]
    text: Optional[str] = None
    level: int = Field(default=2, ge=1, le=6)
    rows: Optional[List[List[str]]] = None
    header_rows: int = Field(default=0, ge=0)
    image: Optional[str] = None


class OutputDocItem(BaseModel):
    title: str = ""
    meta: str = ""
    note: str = ""
    blocks: List[OutputBlock] = []


class OutputDocxRequest(BaseModel):
    title: str = "uSTAT output"
    filename: str = "ustat_output"
    items: List[OutputDocItem]


class OutputDocxError(ValueError):
    """A request this builder refuses; the router turns it into a 422."""


def _image_bytes(data_url: str) -> bytes:
    prefix = next((p for p in _IMAGE_PREFIXES if data_url.startswith(p)), None)
    if prefix is None:
        raise OutputDocxError("Figures must be base64 PNG or JPEG data URLs.")
    try:
        raw = base64.b64decode(data_url[len(prefix):], validate=True)
    except (binascii.Error, ValueError) as exc:
        raise OutputDocxError("A figure is not valid base64.") from exc
    if len(raw) > MAX_IMAGE_BYTES:
        raise OutputDocxError("A figure is larger than 15 MB.")
    if not (raw.startswith(b"\x89PNG\r\n\x1a\n") or raw.startswith(b"\xff\xd8\xff")):
        raise OutputDocxError("A figure is not a PNG or JPEG image.")
    return raw


def _add_table(doc, rows: List[List[str]], header_rows: int) -> None:
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Pt

    n_cols = max((len(r) for r in rows), default=0)
    if n_cols == 0:
        return
    if len(rows) * n_cols > MAX_TABLE_CELLS:
        raise OutputDocxError("A table has more than 20,000 cells.")
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c in range(n_cols):
            cell = table.rows[r].cells[c]
            cell.text = row[c] if c < len(row) else ""
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.bold = r < header_rows
    doc.add_paragraph()


def build_output_docx(req: OutputDocxRequest) -> bytes:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    if not req.items:
        raise OutputDocxError("The output document is empty.")
    if len(req.items) > MAX_ITEMS:
        raise OutputDocxError(f"At most {MAX_ITEMS} items can be exported at once.")

    doc = Document()
    doc.add_heading(req.title, level=0)

    for item in req.items:
        doc.add_heading(item.title or "Result", level=1)
        if item.meta:
            meta = doc.add_paragraph().add_run(item.meta)
            meta.font.size = Pt(8)
            meta.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        if item.note.strip():
            note = doc.add_paragraph().add_run(item.note.strip())
            note.italic = True
        for block in item.blocks:
            if block.type == "heading" and block.text:
                doc.add_heading(block.text, level=min(block.level + 1, 9))
            elif block.type == "paragraph" and block.text:
                doc.add_paragraph(block.text)
            elif block.type == "table" and block.rows:
                _add_table(doc, block.rows, block.header_rows)
            elif block.type == "image" and block.image:
                raw = _image_bytes(block.image)
                try:
                    doc.add_picture(io.BytesIO(raw), width=Inches(_FIGURE_WIDTH_IN))
                except Exception as exc:  # python-docx raises assorted errors on a damaged image
                    raise OutputDocxError("A figure could not be read as an image.") from exc

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
